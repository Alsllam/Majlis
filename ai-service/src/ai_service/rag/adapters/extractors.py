"""
Document extractors. Cloud: Azure Document Intelligence `prebuilt-layout` (PDF, scans, images, Office files) with
Markdown output. `SimpleExtractor` handles digital PDF, DOCX, TXT and MD without any service, which keeps local
development and tests working; the pipeline tries Document Intelligence first when it is configured.
"""

import io
import re

import docx
from azure.ai.documentintelligence.aio import DocumentIntelligenceClient
from azure.ai.documentintelligence.models import AnalyzeDocumentRequest, DocumentContentFormat
from azure.core.credentials import AzureKeyCredential
from azure.identity.aio import DefaultAzureCredential
from pypdf import PdfReader

from ai_service.rag.models import EmptyDocumentError, ExtractedDocument, UnsupportedDocumentError
from ai_service.settings import Settings

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT_TYPES = {"text/plain", "text/markdown"}
DI_TYPES = {
    PDF,
    DOCX,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "image/png",
    "image/jpeg",
}


class SimpleExtractor:
    """Digital files only: no OCR, no layout model."""

    def supports(self, content_type: str) -> bool:
        return content_type in TEXT_TYPES or content_type in {PDF, DOCX}

    async def extract(self, data: bytes, content_type: str, file_name: str) -> ExtractedDocument:
        if content_type in TEXT_TYPES:
            text = data.decode("utf-8", errors="replace")
            return _require_text(text, 1)
        if content_type == PDF:
            return _require_text(*_pdf(data))
        if content_type == DOCX:
            return _require_text(_docx(data), 1)
        raise UnsupportedDocumentError(f"{content_type} ({file_name})")

    async def aclose(self) -> None:
        return None


def _require_text(text: str, pages: int) -> ExtractedDocument:
    if not re.search(r"\w", text):
        raise EmptyDocumentError
    return ExtractedDocument(text, pages)


def _pdf(data: bytes) -> tuple[str, int]:
    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for number, page in enumerate(reader.pages, 1):
        parts.append(f"[[page {number}]]")
        parts.append(page.extract_text() or "")
    return "\n".join(parts), len(reader.pages)


def _docx(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            lines.append("")
            continue
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        if style.lower().startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            lines.append(f"{'#' * min(int(level), 6)} {text}")
        else:
            lines.append(text)
    for table in document.tables:
        lines.append("")
        for row in table.rows:
            lines.append("| " + " | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells) + " |")
    return "\n".join(lines)


class AzureDocumentIntelligenceExtractor:
    """Cloud layout extraction (AI-RAG-001), including scanned Arabic. The only module importing that SDK."""

    def __init__(self, settings: Settings) -> None:
        endpoint = settings.document_intelligence_endpoint
        if settings.document_intelligence_api_key is not None:
            self._client = DocumentIntelligenceClient(
                endpoint, AzureKeyCredential(settings.document_intelligence_api_key.get_secret_value())
            )
        else:
            self._client = DocumentIntelligenceClient(endpoint, DefaultAzureCredential())

    def supports(self, content_type: str) -> bool:
        return content_type in DI_TYPES

    async def extract(self, data: bytes, content_type: str, file_name: str) -> ExtractedDocument:
        poller = await self._client.begin_analyze_document(
            "prebuilt-layout",
            AnalyzeDocumentRequest(bytes_source=data),
            output_content_format=DocumentContentFormat.MARKDOWN,
        )
        result = await poller.result()
        pages = len(result.pages or [])
        text = _page_markers(result.content or "")
        return _require_text(text, max(pages, 1))

    async def aclose(self) -> None:
        await self._client.close()


def _page_markers(markdown: str) -> str:
    """Document Intelligence separates pages with `<!-- PageBreak -->`; the chunker reads `[[page N]]`."""
    parts = re.split(r"<!--\s*PageBreak\s*-->", markdown)
    out: list[str] = []
    for number, part in enumerate(parts, 1):
        out.append(f"[[page {number}]]")
        out.append(re.sub(r"<!--.*?-->", "", part, flags=re.DOTALL))
    return "\n".join(out)
